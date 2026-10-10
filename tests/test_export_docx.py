"""Draft → .docx export — structure, inline formatting, citation
integrity (numbered references resolved through the shared paper lookup),
and the glossary. Round-trips through python-docx, which is itself a
validity check (a corrupt part would fail to re-open)."""

from __future__ import annotations

from pathlib import Path
from typing import Any

import pytest

from precis.dispatch import Hub
from precis.handlers.draft import DraftHandler
from precis.handlers.todo import TodoHandler
from precis.store import ChunkInsert, Store

docx = pytest.importorskip("docx")  # python-docx (the `docx` extra)

from precis.export.docx import export_docx


def _seed_paper(store: Store, slug: str, title: str, year: int) -> None:
    store.insert_ref(kind="paper", slug=slug, title=title, year=year, provider="manual")
    paper_ref = store.get_ref(kind="paper", id=slug)
    assert paper_ref is not None
    store.chunks.insert_chunks(
        paper_ref.id,
        [ChunkInsert(ord=0, text="body", slug="b0")],
    )


@pytest.fixture
def draft(hub: Hub) -> DraftHandler:
    return DraftHandler(hub=hub)


def _make_draft(draft: DraftHandler, hub: Hub) -> object:
    pid = int(
        TodoHandler(hub=hub)
        .put(text="proj")
        .body.split("id=")[1]
        .split()[0]
        .rstrip(",.()")
    )
    draft.put(id="d1", title="CO2 Capture in MOFs", project=pid)
    sec = draft.put(id="d1", chunk_kind="heading", text="Methods", at={"last": True})
    import re

    sec_h = re.search(r"dc\d+", sec.body).group(0)  # type: ignore[union-attr]
    draft.put(
        id="d1",
        chunk_kind="paragraph",
        text="Amine **functionalization** improves uptake [§miller2020~0].",
        at={"into": sec_h, "last": True},
    )
    draft.put(
        id="d1",
        chunk_kind="term",
        text="metal-organic framework",
        meta={"short": "MOF"},
    )
    return hub.live_store.get_ref(kind="draft", id="d1")


def test_export_produces_valid_docx(
    draft: DraftHandler, hub: Hub, tmp_path: Path
) -> None:
    _seed_paper(hub.live_store, "miller2020", "A study of MOFs", 2020)
    ref = _make_draft(draft, hub)
    out = tmp_path / "d1.docx"
    res = export_docx(hub.live_store, ref, target_path=out)
    assert out.is_file()
    # Re-open (validity check) and read the text.
    doc = docx.Document(str(out))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "CO2 Capture in MOFs" in text  # title
    assert "Methods" in text  # heading
    assert "functionalization" in text  # bold run content
    # Citation renders as a numbered superscript marker in the prose, backed
    # by the References section (no native endnote field — those can't repeat).
    assert "[1]" in text
    import zipfile

    with zipfile.ZipFile(out) as z:
        names = set(z.namelist())
        body = z.read("word/document.xml").decode("utf-8")
    assert "word/endnotes.xml" not in names  # no endnote part
    assert "endnoteReference" not in body
    # The [1] mark is a superscript run.
    assert '<w:vertAlign w:val="superscript"/>' in body


def test_citation_integrity_in_references(
    draft: DraftHandler, hub: Hub, tmp_path: Path
) -> None:
    _seed_paper(hub.live_store, "miller2020", "A study of MOFs", 2020)
    ref = _make_draft(draft, hub)
    out = tmp_path / "d1.docx"
    res = export_docx(hub.live_store, ref, target_path=out)
    assert res.cited_slugs == ["miller2020"]
    text = "\n".join(p.text for p in docx.Document(str(out)).paragraphs)
    # A numbered References section carries the entry resolved through the
    # SAME paper lookup as the .bib path (integrity parity with the PDF).
    assert "References" in text
    assert "[1]" in text  # entry number == in-text mark
    assert "A study of MOFs" in text  # resolved title
    assert "2020" in text


def test_repeated_citation_does_not_corrupt(
    draft: DraftHandler, hub: Hub, tmp_path: Path
) -> None:
    """A paper cited at *non-adjacent* sites prints ``[1]`` each time and
    yields ONE References entry — the case a native Word endnote can't model
    (an endnote is 1:1 with its reference, and reusing one makes Word declare
    the document's content unreadable)."""
    _seed_paper(hub.live_store, "wu22", "Wu 2022 study", 2022)
    pid = int(
        TodoHandler(hub=hub)
        .put(text="proj")
        .body.split("id=")[1]
        .split()[0]
        .rstrip(",.()")
    )
    draft.put(id="dr", title="T", project=pid)
    draft.put(
        id="dr",
        chunk_kind="paragraph",
        text="First claim [§wu22~3]. Then unrelated prose. Second claim [§wu22~9].",
        at={"last": True},
    )
    ref = hub.live_store.get_ref(kind="draft", id="dr")
    out = tmp_path / "dr.docx"
    res = export_docx(hub.live_store, ref, target_path=out)
    assert res.cited_slugs == ["wu22"]
    import zipfile

    with zipfile.ZipFile(out) as z:
        body = z.read("word/document.xml").decode("utf-8")
    # Two non-adjacent marks → the same number reused, no endnote machinery.
    assert body.count("endnoteReference") == 0
    text = "\n".join(p.text for p in docx.Document(str(out)).paragraphs)
    assert text.count("[1]") >= 2  # the mark repeats (≥2 sites + References)
    assert "Wu 2022 study" in text  # one resolved entry


def test_docx_embeds_figure_image(
    draft: DraftHandler, hub: Hub, tmp_path: Path
) -> None:
    """A canvas figure rasterises to PNG and embeds as an inline image + an
    italic caption paragraph. Uses a canvas (SVG→PNG via
    resvg) so the embedded raster is one python-docx accepts."""
    pytest.importorskip("resvg_py")
    from precis.handlers.figure import FigureHandler

    pid = int(
        TodoHandler(hub=hub)
        .put(text="proj")
        .body.split("id=")[1]
        .split()[0]
        .rstrip(",.()")
    )
    draft.put(id="d1", title="T", project=pid)
    ref = hub.live_store.get_ref(kind="draft", id="d1")
    assert ref is not None
    title_h = hub.live_store.drafts.reading_order(ref.id)[0].handle
    # a caption-only figure chunk (no blob) backed by a canvas.
    fig = hub.live_store.drafts.add_chunks(
        ref_id=ref.id,
        chunk_kind="figure",
        text="Fig 1. A widget.",
        at={"after": title_h},
        split=False,
    )[0]
    FigureHandler(hub=hub).put(
        id="fig1",
        title="Widget",
        text=(
            '<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 100 60">'
            '<rect x="5" y="5" width="90" height="50" fill="none" stroke="black"/>'
            "</svg>"
        ),
    )
    canvas = hub.live_store.get_ref(kind="figure", id="fig1")
    assert canvas is not None
    hub.live_store.drafts.link_figure_canvas(fig.chunk_id, canvas.id)

    out = tmp_path / "d1.docx"
    export_docx(hub.live_store, ref, target_path=out)

    doc = docx.Document(str(out))
    assert len(doc.inline_shapes) == 1  # the rasterised canvas embedded
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Fig 1. A widget." in text  # caption


def test_docx_withheld_figure_is_boxed_not_embedded(
    draft: DraftHandler, hub: Hub, tmp_path: Path
) -> None:
    """A handle in ``withheld_figures`` renders a bordered notice naming the
    publisher and source plus the caption; the image is never embedded."""
    import base64

    pid = int(
        TodoHandler(hub=hub)
        .put(text="proj")
        .body.split("id=")[1]
        .split()[0]
        .rstrip(",.()")
    )
    draft.put(id="d1", title="T", project=pid)
    ref = hub.live_store.get_ref(kind="draft", id="d1")
    assert ref is not None
    title_h = hub.live_store.drafts.reading_order(ref.id)[0].handle
    # 1x1 PNG
    png = base64.b64encode(
        base64.b64decode(
            "iVBORw0KGgoAAAANSUhE"
            + "UgAAAAEAAAABCAYAAAAf"
            + "FcSJAAAADUlEQVR42mNk"
            + "YPhfDwAChwGA60e6kgAA"
            + "AABJRU5ErkJggg=="
        )
    ).decode()
    draft.put(
        id="d1",
        chunk_kind="figure",
        text="Fig 1. Borrowed.",
        image=png,
        origin="third_party",
        permission={
            "publisher": "ACME",
            "source_paper": "Doe 2020",
            "status": "requested",
        },
        at={"after": f"\u00b6{title_h}"},
    )
    figc = next(
        c
        for c in hub.live_store.drafts.reading_order(ref.id)
        if c.chunk_kind == "figure"
    )
    out = tmp_path / "d1.docx"
    export_docx(
        hub.live_store, ref, target_path=out, withheld_figures=frozenset({figc.dc})
    )
    doc = docx.Document(str(out))
    assert len(doc.inline_shapes) == 0
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Figure withheld pending permission: ACME, Doe 2020" in text
    assert "Fig 1. Borrowed." in text


def test_glossary_terms_become_a_sorted_used_only_acronym_list(
    draft: DraftHandler, hub: Hub, tmp_path: Path
) -> None:
    """The draft's own Glossary heading + term chunks are not body; the back
    matter is an "Acronyms" list of the shorts the prose used, sorted
    case-insensitively (LaTeX's generated acronym list)."""
    pid = _new_draft_project(hub)
    draft.put(id="dgl", title="T", project=pid)
    for short, long in [
        ("pG", "pristine graphene"),
        ("CNT", "carbon nanotube"),
        ("DFT", "density functional theory"),  # defined, never used
        ("ha", "hexylamine"),
    ]:
        draft.put(id="dgl", chunk_kind="term", text=long, meta={"short": short})
    draft.put(
        id="dgl",
        chunk_kind="paragraph",
        text="We study pG, ha and CNT here.",
        at={"last": True},
    )
    ref = hub.live_store.get_ref(kind="draft", id="dgl")
    out = tmp_path / "dgl.docx"
    export_docx(hub.live_store, ref, target_path=out)
    texts = [p.text for p in docx.Document(str(out)).paragraphs]
    assert "Glossary" not in texts
    i = texts.index("Acronyms")
    assert texts[i + 1 :] == [
        "CNT — carbon nanotube",
        "ha — hexylamine",
        "pG — pristine graphene",
    ]
    assert not any(t.startswith("DFT") for t in texts)  # never used in the body


def test_missing_paper_warns_but_exports(
    draft: DraftHandler, hub: Hub, tmp_path: Path
) -> None:
    ref = _make_draft(draft, hub)  # cites miller2020 which is NOT seeded
    out = tmp_path / "d1.docx"
    res = export_docx(hub.live_store, ref, target_path=out)
    assert out.is_file()
    assert any("miller2020" in w for w in res.warnings)


def test_acronym_first_use_expansion(
    draft: DraftHandler, hub: Hub, tmp_path: Path
) -> None:
    pid = int(
        TodoHandler(hub=hub)
        .put(text="proj")
        .body.split("id=")[1]
        .split()[0]
        .rstrip(",.()")
    )
    draft.put(id="dx", title="T", project=pid)
    draft.put(
        id="dx",
        chunk_kind="term",
        text="metal-organic framework",
        meta={"short": "MOF"},
    )
    draft.put(
        id="dx",
        chunk_kind="paragraph",
        text="First, MOF systems. Later, more MOFs appear.",
        at={"last": True},
    )
    ref = hub.live_store.get_ref(kind="draft", id="dx")
    out = tmp_path / "dx.docx"
    export_docx(hub.live_store, ref, target_path=out)
    text = "\n".join(p.text for p in docx.Document(str(out)).paragraphs)
    # First prose occurrence expanded; later plural stays abbreviated.
    assert "metal-organic framework (MOF)" in text
    assert "MOFs appear" in text  # plural, not expanded
    # The used abbreviation lands in the generated Acronyms list.
    assert "Acronyms" in text
    assert "MOF — metal-organic framework" in text
    assert "Glossary" not in text


def test_math_renders_as_omml(draft: DraftHandler, hub: Hub, tmp_path: Path) -> None:
    pytest.importorskip("latex2mathml")
    pid = int(
        TodoHandler(hub=hub)
        .put(text="proj")
        .body.split("id=")[1]
        .split()[0]
        .rstrip(",.()")
    )
    draft.put(id="dm", title="T", project=pid)
    draft.put(
        id="dm",
        chunk_kind="paragraph",
        text="The relation $E = mc^2$ and a fraction $\\frac{a}{b}$.",
        at={"last": True},
    )
    ref = hub.live_store.get_ref(kind="draft", id="dm")
    out = tmp_path / "dm.docx"
    export_docx(hub.live_store, ref, target_path=out)
    # Inspect the document XML for native OMML math.
    import zipfile

    with zipfile.ZipFile(out) as z:
        doc_xml = z.read("word/document.xml").decode("utf-8")
    assert "oMath" in doc_xml  # native math, not literal "$...$"
    assert "sSup" in doc_xml  # the c^2 superscript
    assert "}f" in doc_xml or "<m:f" in doc_xml or ":f>" in doc_xml  # the fraction
    assert "$" not in doc_xml  # no leftover literal math source
    # Re-open as a validity check.
    docx.Document(str(out))


def test_garbled_and_money_math_demoted_to_literal_text(
    draft: DraftHandler, hub: Hub, tmp_path: Path
) -> None:
    r"""Parity with the LaTeX exporter's demotion predicates
    (``_math_braces_balanced`` / ``_math_plausible``, shared via import):
    an unbalanced-brace span and a money-dollar mispairing render as
    literal prose runs, never OMML; an author-escaped ``\$`` renders as
    the bare ``$`` (Word has no escape syntax)."""
    pytest.importorskip("latex2mathml")
    pid = int(
        TodoHandler(hub=hub)
        .put(text="proj")
        .body.split("id=")[1]
        .split()[0]
        .rstrip(",.()")
    )
    draft.put(id="dg", title="T", project=pid)
    draft.put(
        id="dg",
        chunk_kind="paragraph",
        text=(
            r"The ratio $\sqrt{2/\sqrt{3}$ holds. Oligomers cost "
            r"$10-50 each, versus $200 for staples, plus \$99 fees."
        ),
        at={"last": True},
    )
    ref = hub.live_store.get_ref(kind="draft", id="dg")
    out = tmp_path / "dg.docx"
    export_docx(hub.live_store, ref, target_path=out)
    import zipfile

    with zipfile.ZipFile(out) as z:
        doc_xml = z.read("word/document.xml").decode("utf-8")
    assert "oMath" not in doc_xml  # nothing here is math
    text = "\n".join(p.text for p in docx.Document(str(out)).paragraphs)
    assert r"$\sqrt{2/\sqrt{3}$" in text  # demoted verbatim, not OMML
    assert "$10-50 each, versus $200 for staples" in text
    assert "$99 fees" in text and r"\$" not in text  # \$ → bare $


def test_empty_base_math_gets_a_base(
    draft: DraftHandler, hub: Hub, tmp_path: Path
) -> None:
    """``Zr$_6$`` / ``UO$_2^{2+}$`` put the base outside the math, leaving an
    empty subscript base — an empty OMML ``<m:e/>`` Word draws as a dotted-box
    placeholder. The exporter folds the adjacent token into the math instead."""
    pytest.importorskip("latex2mathml")
    pid = int(
        TodoHandler(hub=hub)
        .put(text="proj")
        .body.split("id=")[1]
        .split()[0]
        .rstrip(",.()")
    )
    draft.put(id="dz", title="T", project=pid)
    draft.put(
        id="dz",
        chunk_kind="paragraph",
        text="The Zr$_6$ node and the UO$_2^{2+}$ ion.",
        at={"last": True},
    )
    ref = hub.live_store.get_ref(kind="draft", id="dz")
    out = tmp_path / "dz.docx"
    export_docx(hub.live_store, ref, target_path=out)
    import zipfile

    with zipfile.ZipFile(out) as z:
        doc_xml = z.read("word/document.xml").decode("utf-8")
    assert "<m:e/>" not in doc_xml  # no empty base → no dotted box
    assert "oMath" in doc_xml
    docx.Document(str(out))


def test_standalone_equation_numbered_and_cross_referenced(
    draft: DraftHandler, hub: Hub, tmp_path: Path
) -> None:
    """A paragraph chunk whose ENTIRE text is one ``$$…$$`` span numbers
    "(N)" next to the native OMML math, and a bare ``[dc<id>]`` cross-ref
    to that chunk auto-resolves to "Eq. (N)" — no new authoring syntax.
    docx has no live cross-reference field (unlike the PDF path's
    cleveref), so the number is a static count resolved at export time."""
    pytest.importorskip("latex2mathml")
    import re

    pid = int(
        TodoHandler(hub=hub)
        .put(text="proj")
        .body.split("id=")[1]
        .split()[0]
        .rstrip(",.()")
    )
    draft.put(id="eqd", title="T", project=pid)
    eq1 = draft.put(
        id="eqd", chunk_kind="paragraph", text="$$E = mc^2$$", at={"last": True}
    )
    dc1 = re.search(r"dc\d+", eq1.body).group(0)  # type: ignore[union-attr]
    draft.put(
        id="eqd",
        chunk_kind="paragraph",
        text=f"As shown in [{dc1}], mass and energy relate.",
        at={"last": True},
    )
    draft.put(id="eqd", chunk_kind="paragraph", text="$$F = ma$$", at={"last": True})

    ref = hub.live_store.get_ref(kind="draft", id="eqd")
    out = tmp_path / "eqd.docx"
    export_docx(hub.live_store, ref, target_path=out)
    text = "\n".join(p.text for p in docx.Document(str(out)).paragraphs)
    assert "(1)" in text and "(2)" in text  # the two equations, numbered in order
    assert "Eq. (1)" in text  # the bare [dc<id>] cross-ref auto-resolved
    import zipfile

    with zipfile.ZipFile(out) as z:
        doc_xml = z.read("word/document.xml").decode("utf-8")
    assert doc_xml.count("oMath") >= 4  # two equations, each open+close tag


def test_starred_equation_is_unnumbered_in_docx(
    draft: DraftHandler, hub: Hub, tmp_path: Path
) -> None:
    """A trailing ``*`` right after the closing ``$$`` opts a display
    equation out of numbering — no "(N)" label, and the marker itself
    never leaks into the rendered text."""
    pytest.importorskip("latex2mathml")
    pid = int(
        TodoHandler(hub=hub)
        .put(text="proj")
        .body.split("id=")[1]
        .split()[0]
        .rstrip(",.()")
    )
    draft.put(id="eqs", title="T", project=pid)
    draft.put(
        id="eqs",
        chunk_kind="paragraph",
        text="$$a^2 + b^2 = c^2$$*",
        at={"last": True},
    )
    ref = hub.live_store.get_ref(kind="draft", id="eqs")
    out = tmp_path / "eqs.docx"
    export_docx(hub.live_store, ref, target_path=out)
    text = "\n".join(p.text for p in docx.Document(str(out)).paragraphs)
    assert "(1)" not in text
    assert "*" not in text  # the star marker itself never leaks into output
    import zipfile

    with zipfile.ZipFile(out) as z:
        doc_xml = z.read("word/document.xml").decode("utf-8")
    assert "oMath" in doc_xml


def test_latex_cite_command_is_folded(
    draft: DraftHandler, hub: Hub, tmp_path: Path
) -> None:
    """A draft carrying verbatim LaTeX ``\\cite{key}`` resolves to a numbered
    mark + References entry — the ``\\cite{`` / ``}`` wrapper never leaks as
    literal text."""
    _seed_paper(hub.live_store, "wu22", "Wu 2022 study", 2022)
    pid = int(
        TodoHandler(hub=hub)
        .put(text="proj")
        .body.split("id=")[1]
        .split()[0]
        .rstrip(",.()")
    )
    draft.put(id="dc", title="T", project=pid)
    draft.put(
        id="dc",
        chunk_kind="paragraph",
        text="Adsorption is fast \\cite{wu22}.",
        at={"last": True},
    )
    ref = hub.live_store.get_ref(kind="draft", id="dc")
    out = tmp_path / "dc.docx"
    res = export_docx(hub.live_store, ref, target_path=out)
    assert res.cited_slugs == ["wu22"]
    text = "\n".join(p.text for p in docx.Document(str(out)).paragraphs)
    assert "\\cite" not in text and "cite{" not in text  # wrapper folded away
    assert "[1]" in text  # resolved to a numbered mark
    assert "Wu 2022 study" in text  # References entry


def test_handle_form_citation_resolves(
    draft: DraftHandler, hub: Hub, tmp_path: Path
) -> None:
    """A draft that cites a paper by universal handle (``[pa<ref_id>]``, the
    form every LaTeX-imported draft uses) must produce a numbered mark +
    References entry — NOT render as nothing. Regression: the docx exporter
    used to drop handle citations entirely (the LaTeX/PDF path resolved them),
    so handle-cited drafts exported with no citations and no References
    section at all."""
    from precis.utils import handle_registry

    _seed_paper(hub.live_store, "nasibulin2007", "Multifunctional nanobuds", 2007)
    pref = hub.live_store.get_ref(kind="paper", id="nasibulin2007")
    assert pref is not None
    handle = handle_registry.format_handle("paper", pref.id)  # 'pa<ref_id>'
    assert handle.startswith("pa")
    pid = int(
        TodoHandler(hub=hub)
        .put(text="proj")
        .body.split("id=")[1]
        .split()[0]
        .rstrip(",.()")
    )
    draft.put(id="dh", title="T", project=pid)
    draft.put(
        id="dh",
        chunk_kind="paragraph",
        text=f"Nanobuds were first reported [{handle}].",
        at={"last": True},
    )
    ref = hub.live_store.get_ref(kind="draft", id="dh")
    out = tmp_path / "dh.docx"
    res = export_docx(hub.live_store, ref, target_path=out)
    assert res.cited_slugs == ["nasibulin2007"]  # handle resolved to the slug
    text = "\n".join(p.text for p in docx.Document(str(out)).paragraphs)
    assert "[1]" in text  # numbered mark emitted (not dropped)
    assert "References" in text
    assert "Multifunctional nanobuds" in text  # resolved entry


def test_computed_evidence_handle_renders_as_text(
    draft: DraftHandler, hub: Hub, tmp_path: Path
) -> None:
    """A bracket-handle cite of a computational-evidence kind (a simulation
    structure, ``[st<id>]``) is real grounding, not a bibliography entry —
    it must render as plain text (handle, or display-form surface text)
    rather than being silently dropped."""
    from precis.utils import handle_registry

    art = hub.live_store.insert_ref(
        kind="structure", slug="pd-relaxed", title="Pd relaxed cell", meta={}
    )
    handle = handle_registry.format_handle("structure", art.id)  # 'st<ref_id>'
    assert handle.startswith("st")
    pid = int(
        TodoHandler(hub=hub)
        .put(text="proj")
        .body.split("id=")[1]
        .split()[0]
        .rstrip(",.()")
    )
    draft.put(id="dc-evi", title="T", project=pid)
    draft.put(
        id="dc-evi",
        chunk_kind="paragraph",
        text=(
            f"The relaxed geometry [{handle}] and the "
            f"[computed relaxation]({handle}) confirm the bond length."
        ),
        at={"last": True},
    )
    ref = hub.live_store.get_ref(kind="draft", id="dc-evi")
    out = tmp_path / "dc-evi.docx"
    export_docx(hub.live_store, ref, target_path=out)
    text = "\n".join(p.text for p in docx.Document(str(out)).paragraphs)
    assert handle in text  # bare bracket handle rendered as plain text
    assert "computed relaxation" in text  # display-form surface text kept


def test_endnote_citations_emit_cwyw_fields(
    draft: DraftHandler, hub: Hub, tmp_path: Path
) -> None:
    """``citations='endnote'`` emits native EndNote Cite-While-You-Write
    fields (``ADDIN EN.CITE`` carrying the full <record>, plus an
    ``EN.REFLIST`` bibliography field and the EN.* doc-vars) instead of a
    plain ``[n]`` marker — the format EndNote recognizes and can reformat.
    Structure pinned against a real EndNote-authored .docx."""
    import zipfile

    _seed_paper(hub.live_store, "nasibulin2007", "Multifunctional nanobuds", 2007)
    pid = int(
        TodoHandler(hub=hub)
        .put(text="proj")
        .body.split("id=")[1]
        .split()[0]
        .rstrip(",.()")
    )
    draft.put(id="en", title="T", project=pid)
    draft.put(
        id="en",
        chunk_kind="paragraph",
        text="Nanobuds were first reported [§nasibulin2007~0].",
        at={"last": True},
    )
    ref = hub.live_store.get_ref(kind="draft", id="en")
    out = tmp_path / "en.docx"
    res = export_docx(hub.live_store, ref, target_path=out, citations="endnote")
    assert res.cited_slugs == ["nasibulin2007"]
    with zipfile.ZipFile(out) as z:
        body = z.read("word/document.xml").decode("utf-8")
        settings = z.read("word/settings.xml").decode("utf-8")
    # In-text EN.CITE complex field carrying the traveling-library record.
    assert " ADDIN EN.CITE &lt;EndNote&gt;&lt;Cite&gt;" in body
    assert '&lt;ref-type name="Journal Article"&gt;17' in body
    assert "Multifunctional nanobuds" in body  # embedded record title
    assert '<w:fldChar w:fldCharType="begin"/>' in body
    # Bibliography field + the document-level EndNote state EndNote reads.
    assert " ADDIN EN.REFLIST " in body
    assert "EN.InstantFormat" in settings
    assert "EN.Layout" in settings
    assert "EN.Libraries" in settings
    # Re-open as a validity check (a malformed part would fail here).
    docx.Document(str(out))
    # Plain mode is unchanged — no EndNote machinery leaks in.
    plain = tmp_path / "plain.docx"
    export_docx(hub.live_store, ref, target_path=plain, citations="plain")
    with zipfile.ZipFile(plain) as z:
        pbody = z.read("word/document.xml").decode("utf-8")
    assert "ADDIN EN.CITE" not in pbody
    assert "[1]" in "\n".join(p.text for p in docx.Document(str(plain)).paragraphs)


def test_endnote_pc_citation_embeds_cited_passage(
    draft: DraftHandler, hub: Hub, tmp_path: Path
) -> None:
    """A ``[pc<id>]`` chunk citation embeds that chunk's text as the EndNote
    record's ``<research-notes>`` (the traveling note) — so the exact cited
    passage rides along. A ``[pa<id>]`` ref-level cite (whole paper) carries
    no passage."""
    import zipfile

    store = hub.live_store
    store.insert_ref(
        kind="paper", slug="nas07", title="Nanobuds paper", year=2007, provider="manual"
    )
    pref = store.get_ref(kind="paper", id="nas07")
    assert pref is not None
    passage = "The exact cited passage about sp3 rehybridization at junctions."
    store.chunks.insert_chunks(pref.id, [ChunkInsert(ord=0, text=passage, slug="b0")])
    with store.pool.connection() as conn:
        row = conn.execute(
            "SELECT chunk_id FROM chunks WHERE ref_id = %s AND ord >= 0 "
            "ORDER BY ord LIMIT 1",
            (pref.id,),
        ).fetchone()
        assert row is not None
        chunk_id = int(row[0])
    pid = int(
        TodoHandler(hub=hub)
        .put(text="proj")
        .body.split("id=")[1]
        .split()[0]
        .rstrip(",.()")
    )
    draft.put(id="pc", title="T", project=pid)
    draft.put(
        id="pc",
        chunk_kind="paragraph",
        text=f"Junctions rehybridize [pc{chunk_id}].",
        at={"last": True},
    )
    ref = store.get_ref(kind="draft", id="pc")
    out = tmp_path / "pc.docx"
    res = export_docx(store, ref, target_path=out, citations="endnote")
    assert res.cited_slugs == ["nas07"]
    body = zipfile.ZipFile(out).read("word/document.xml").decode("utf-8")
    # The cited chunk's text is embedded as the record's Research Notes.
    assert "&lt;research-notes&gt;" in body
    assert "sp3 rehybridization at junctions" in body  # the passage rode along
    docx.Document(str(out))  # validity


def test_omml_converter_returns_none_on_empty() -> None:
    pytest.importorskip("latex2mathml")
    from precis.export.omml import latex_to_omml

    assert latex_to_omml("") is None
    assert latex_to_omml("   ") is None
    el = latex_to_omml("x^2")  # well-formed → an <m:oMath> element
    assert el is not None and el.tag.endswith("oMath")


def test_same_paper_chunks_collapse_to_one_mark(
    draft: DraftHandler, hub: Hub, tmp_path: Path
) -> None:
    """a~3, a~9, a~23 (different chunks, same paper) → ONE References entry,
    and consecutive marks collapse to a single ``[1]``."""
    _seed_paper(hub.live_store, "wu22", "Wu 2022 study", 2022)
    pid = int(
        TodoHandler(hub=hub)
        .put(text="proj")
        .body.split("id=")[1]
        .split()[0]
        .rstrip(",.()")
    )
    draft.put(id="dd", title="T", project=pid)
    draft.put(
        id="dd",
        chunk_kind="paragraph",
        text="Several findings [§wu22~3] [§wu22~9] [§wu22~23] agree.",
        at={"last": True},
    )
    ref = hub.live_store.get_ref(kind="draft", id="dd")
    out = tmp_path / "dd.docx"
    res = export_docx(hub.live_store, ref, target_path=out)
    assert res.cited_slugs == ["wu22"]  # one paper, deduped
    paras = [p.text for p in docx.Document(str(out)).paragraphs]
    body_para = next(p for p in paras if "Several findings" in p)
    assert body_para.count("[1]") == 1  # consecutive marks collapsed to one
    assert "Wu 2022 study" in "\n".join(paras)  # one References entry


def test_export_renders_list_styles(
    draft: DraftHandler, hub: Hub, tmp_path: Path
) -> None:
    """ulist/olist items get Word's built-in List Bullet / List Number
    styles; the container itself emits no paragraph (migration 0037)."""
    import re

    pid = int(
        TodoHandler(hub=hub)
        .put(text="proj")
        .body.split("id=")[1]
        .split()[0]
        .rstrip(",.()")
    )
    draft.put(id="ld", title="Lists", project=pid)
    ul = draft.put(id="ld", chunk_kind="ulist", text="list", at={"last": True})
    ul_h = re.search(r"dc\d+", ul.body).group(0)  # type: ignore[union-attr]
    draft.put(id="ld", chunk_kind="item", text="alpha", at={"into": ul_h, "last": True})
    ol = draft.put(id="ld", chunk_kind="olist", text="list", at={"last": True})
    ol_h = re.search(r"dc\d+", ol.body).group(0)  # type: ignore[union-attr]
    draft.put(id="ld", chunk_kind="item", text="one", at={"into": ol_h, "last": True})

    ref = hub.live_store.get_ref(kind="draft", id="ld")
    out = tmp_path / "ld.docx"
    export_docx(hub.live_store, ref, target_path=out)
    doc = docx.Document(str(out))
    styled = {
        p.text: p.style.name for p in doc.paragraphs if p.text in ("alpha", "one")
    }
    assert styled.get("alpha") == "List Bullet"
    assert styled.get("one") == "List Number"


def test_markdown_bullets_export_as_nested_list_styles(
    draft: DraftHandler, hub: Hub, tmp_path: Path
) -> None:
    """End to end, the trap this closes: bullet text written into a
    paragraph used to reach Word as one run-on Body Text line. It now
    lands structured (:mod:`precis.draft.mdlist`), so the nesting level
    picks List Bullet 2."""
    pid = int(
        TodoHandler(hub=hub)
        .put(text="proj")
        .body.split("id=")[1]
        .split()[0]
        .rstrip(",.()")
    )
    draft.put(id="md", title="Lists", project=pid)
    draft.put(
        id="md",
        chunk_kind="paragraph",
        text="- NO side\n    - Bader charge shows 0.39 e\n- NH3 side",
        at={"last": True},
    )

    ref = hub.live_store.get_ref(kind="draft", id="md")
    out = tmp_path / "md.docx"
    export_docx(hub.live_store, ref, target_path=out)
    doc = docx.Document(str(out))
    styled = {p.text: p.style.name for p in doc.paragraphs if p.text}
    assert styled.get("NO side") == "List Bullet"
    assert styled.get("NH3 side") == "List Bullet"
    assert styled.get("Bader charge shows 0.39 e") == "List Bullet 2"


def test_export_renders_table(draft: DraftHandler, hub: Hub, tmp_path: Path) -> None:
    """A chunk_kind='table' becomes a native Word table:
    header row bold, one body row per data row, cells via the inline grammar.
    The derived pipe markdown is not dumped as a paragraph."""
    pid = int(
        TodoHandler(hub=hub)
        .put(text="proj")
        .body.split("id=")[1]
        .split()[0]
        .rstrip(",.()")
    )
    draft.put(id="tb", title="T", project=pid)
    draft.put(
        id="tb",
        chunk_kind="table",
        table={"header": ["ID", "Title"], "rows": [["I1", "alpha"], ["I2", "beta"]]},
        caption="Issue register",
        at={"last": True},
    )
    ref = hub.live_store.get_ref(kind="draft", id="tb")
    out = tmp_path / "tb.docx"
    export_docx(hub.live_store, ref, target_path=out)

    doc = docx.Document(str(out))
    assert len(doc.tables) == 1
    t = doc.tables[0]
    assert len(t.rows) == 3 and len(t.columns) == 2  # header + 2 body rows
    assert [c.text for c in t.rows[0].cells] == ["ID", "Title"]
    assert [c.text for c in t.rows[1].cells] == ["I1", "alpha"]
    assert t.rows[0].cells[0].paragraphs[0].runs[0].bold  # header bold
    # caption rendered as a bold lead-in paragraph; pipe markdown not dumped
    paras = [p.text for p in doc.paragraphs]
    assert "Issue register" in paras
    assert not any("| ID | Title |" in p for p in paras)


def test_render_byline_names_marks_and_ror_link() -> None:
    """The byline block: names with superscript marks (when >1 affiliation)
    + one affiliation paragraph each, ROR org rendered as a real hyperlink.
    Pure over a python-docx Document (no store)."""
    from precis.export.docx import _render_byline
    from precis.utils.authors import build_byline

    doc = docx.Document()
    doc.add_heading("Title", level=0)
    byline = build_byline(
        [
            {"name": "Doe, Jane", "affiliation": "MIT", "ror": "https://ror.org/x"},
            {"name": "Roe, John", "affiliation": "Caltech"},
        ]
    )
    _render_byline(doc, byline)
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Jane Doe" in text and "John Roe" in text
    # superscript marks present as run text on the names paragraph
    names_p = doc.paragraphs[1]
    assert any(r.font.superscript and r.text in ("1", "2") for r in names_p.runs)
    # the ROR affiliation is a real external hyperlink relationship
    xml = doc.paragraphs[2]._p.xml
    assert "hyperlink" in xml
    rels = doc.part.rels
    assert any(r.reltype.endswith("hyperlink") for r in rels.values())


def test_render_byline_orcid_is_a_superscript_id_link() -> None:
    """An author's ORCID renders as a superscript ``iD`` hyperlink to the
    orcid.org record, right after the name."""
    from precis.export.docx import _render_byline
    from precis.utils.authors import build_byline

    doc = docx.Document()
    doc.add_heading("Title", level=0)
    _render_byline(
        doc,
        build_byline(
            [
                {"name": "Doe, Jane", "orcid": "0000-0002-1825-0097"},
                {"name": "Roe, John"},
            ]
        ),
    )
    names_p = doc.paragraphs[1]
    xml = names_p._p.xml
    assert "hyperlink" in xml and ">iD<" in xml
    assert 'w:vertAlign w:val="superscript"' in xml
    rels = doc.part.rels
    assert any(
        r.reltype.endswith("hyperlink")
        and r.target_ref == "https://orcid.org/0000-0002-1825-0097"
        for r in rels.values()
    )


def test_render_byline_empty_authors_is_noop() -> None:
    from precis.export.docx import _render_byline
    from precis.utils.authors import build_byline

    doc = docx.Document()
    doc.add_heading("Title", level=0)
    before = len(doc.paragraphs)
    _render_byline(doc, build_byline(None))
    assert len(doc.paragraphs) == before


class _RefStore:
    """Minimal store for the pure reference/EndNote resolvers: resolves a
    slug to a paper/patent/datasheet ref, no DOI/arXiv aliases."""

    #: Set by the SI tests to a stub whose ``connection()`` yields a dummy.
    pool: Any = None

    def __init__(self, refs):
        self._refs = refs  # (kind, slug) -> Ref-ish

    def get_ref(self, *, kind, id):
        return self._refs.get((kind, id))

    def identifiers_for_refs(self, ref_ids):
        return {}


def test_format_reference_resolves_datasheet_not_stub() -> None:
    """A cited datasheet resolves to a real reference line — parity with the
    .bib path (gr52396) — not the 'missing source' stub."""
    from types import SimpleNamespace

    from precis.export.docx import _format_reference

    store = _RefStore(
        {
            ("datasheet", "stm32f4"): SimpleNamespace(
                id=7,
                slug="stm32f4",
                kind="datasheet",
                title="STM32F4 Reference Manual",
                authors=[{"name": "STMicroelectronics"}],
                year=2019,
                meta={},
            )
        }
    )
    warnings: list[str] = []
    line = _format_reference(store, "stm32f4", warnings)
    assert "STM32F4 Reference Manual" in line
    assert "STMicroelectronics" in line
    assert "[Datasheet]" in line  # sub-type genre label
    assert "missing source" not in line


def test_format_reference_datasheet_uses_vendor_subtype_part() -> None:
    """A bylineless datasheet: vendor stands in as the org, the sub-type
    label + documented part ride along — parity with the .bib @manual entry."""
    from types import SimpleNamespace

    from precis.export.docx import _format_reference

    store = _RefStore(
        {
            ("datasheet", "esp32c3"): SimpleNamespace(
                id=9,
                slug="esp32c3",
                kind="datasheet",
                title="ESP32-C3 App Note",
                authors=None,
                year=2022,
                meta={
                    "vendor": "Espressif Systems",
                    "subtype": "app-note",
                    "part_lcsc": "C2934569",
                },
            )
        }
    )
    warnings: list[str] = []
    line = _format_reference(store, "esp32c3", warnings)
    assert "Espressif Systems" in line
    assert "(2022)" in line
    assert "[Application note]" in line
    assert "Part C2934569" in line
    assert warnings == []


# ── Taproot claim-hub finding handle (Phase 1 — living citations reach
# draft export): a [fi<id>] finding handle that resolves to a
# TAPROOT:claim hub cites its *derived* establishes originator(s) via the
# ONE shared resolver (precis.taproot.cite), not a stored
# primary_cite_key. Mirrors tests/test_export_latex.py's hub coverage.


def _mint_hub_claim(store: Store) -> int:
    from precis.taproot.canon import CanonicalClaim
    from precis.taproot.hub import mint_hub

    claim = CanonicalClaim(
        sentence="Pd/C catalyzes Suzuki coupling at room temperature with a mild base.",
        scope={"material": "Pd/C", "method": "Suzuki coupling", "regime": "RT"},
    )
    return mint_hub(store, claim)


def _new_draft_project(hub: Hub) -> int:
    return int(
        TodoHandler(hub=hub)
        .put(text="proj")
        .body.split("id=")[1]
        .split()[0]
        .rstrip(",.()")
    )


def test_hub_finding_single_originator_renders_one_mark(
    draft: DraftHandler, hub: Hub, tmp_path: Path
) -> None:
    from precis.taproot.hub import attach_evidence
    from precis.utils import handle_registry

    hub_ref = _mint_hub_claim(hub.live_store)
    origin = hub.live_store.insert_ref(
        kind="paper", slug="docxo01", title="Original report", year=2001, meta={}
    ).id
    follow = hub.live_store.insert_ref(
        kind="paper", slug="docxf05", title="Follow-up", year=2005, meta={}
    ).id
    attach_evidence(
        hub.live_store, hub_ref_id=hub_ref, paper_ref_id=origin, role="corroborates"
    )
    attach_evidence(
        hub.live_store, hub_ref_id=hub_ref, paper_ref_id=follow, role="corroborates"
    )
    hub.live_store.add_link(src_ref_id=follow, dst_ref_id=origin, relation="cites")
    finding_handle = handle_registry.format_handle("finding", hub_ref)

    pid = _new_draft_project(hub)
    draft.put(id="dhub1", title="T", project=pid)
    draft.put(
        id="dhub1",
        chunk_kind="paragraph",
        text=f"Living citation [{finding_handle}].",
        at={"last": True},
    )
    ref = hub.live_store.get_ref(kind="draft", id="dhub1")
    out = tmp_path / "dhub1.docx"
    res = export_docx(hub.live_store, ref, target_path=out)

    assert res.cited_slugs == ["docxo01"]  # the derived originator, not corroborator
    text = "\n".join(p.text for p in docx.Document(str(out)).paragraphs)
    assert "[1]" in text
    assert "Original report" in text


def test_hub_finding_multiple_originators_renders_two_marks(
    draft: DraftHandler, hub: Hub, tmp_path: Path
) -> None:
    """docx has no ``\\cite{a,b}`` literal — a multi-originator hub emits
    one numbered mark PER key, per-key calls into the existing ``_cite``
    path (mirrors the LaTeX exporter's grouped-\\cite counterpart)."""
    from precis.taproot.hub import attach_evidence
    from precis.utils import handle_registry

    hub_ref = _mint_hub_claim(hub.live_store)
    a = hub.live_store.insert_ref(
        kind="paper", slug="docxa01", title="A — first report", year=2001, meta={}
    ).id
    b = hub.live_store.insert_ref(
        kind="paper", slug="docxb02", title="B — second report", year=2002, meta={}
    ).id
    citer = hub.live_store.insert_ref(
        kind="paper", slug="docxc09", title="Citer", year=2009, meta={}
    ).id
    for p in (a, b, citer):
        attach_evidence(
            hub.live_store, hub_ref_id=hub_ref, paper_ref_id=p, role="corroborates"
        )
    hub.live_store.add_link(src_ref_id=citer, dst_ref_id=a, relation="cites")
    hub.live_store.add_link(src_ref_id=citer, dst_ref_id=b, relation="cites")
    finding_handle = handle_registry.format_handle("finding", hub_ref)

    pid = _new_draft_project(hub)
    draft.put(id="dhub2", title="T", project=pid)
    draft.put(
        id="dhub2",
        chunk_kind="paragraph",
        text=f"Living citation [{finding_handle}].",
        at={"last": True},
    )
    ref = hub.live_store.get_ref(kind="draft", id="dhub2")
    out = tmp_path / "dhub2.docx"
    res = export_docx(hub.live_store, ref, target_path=out)

    assert res.cited_slugs == ["docxa01", "docxb02"]
    text = "\n".join(p.text for p in docx.Document(str(out)).paragraphs)
    assert "[1]" in text and "[2]" in text
    assert "A — first report" in text and "B — second report" in text


def test_hub_finding_multi_cite_gets_doi_ul_hyperlink_runs(
    draft: DraftHandler, hub: Hub, tmp_path: Path
) -> None:
    """Cite-links: each numbered mark gets a small doi/UL hyperlink-run
    pair right after it (mirrors the LaTeX exporter's ``\\href`` pair),
    and none at all with both ``doi_links=False`` and
    ``library_links=False``."""
    from precis.taproot.hub import attach_evidence
    from precis.utils import handle_registry

    hub_ref = _mint_hub_claim(hub.live_store)
    a = hub.live_store.insert_ref(
        kind="paper", slug="docxa21", title="A — first report", year=2001, meta={}
    ).id
    b = hub.live_store.insert_ref(
        kind="paper", slug="docxb22", title="B — second report", year=2002, meta={}
    ).id
    citer = hub.live_store.insert_ref(
        kind="paper", slug="docxc29", title="Citer", year=2009, meta={}
    ).id
    for p in (a, b, citer):
        attach_evidence(
            hub.live_store, hub_ref_id=hub_ref, paper_ref_id=p, role="corroborates"
        )
    hub.live_store.add_link(src_ref_id=citer, dst_ref_id=a, relation="cites")
    hub.live_store.add_link(src_ref_id=citer, dst_ref_id=b, relation="cites")
    hub.live_store.insert_ref_identifiers(a, [("doi", "10.1/docxa21", "manual")])
    hub.live_store.insert_ref_identifiers(b, [("doi", "10.1/docxb22", "manual")])
    finding_handle = handle_registry.format_handle("finding", hub_ref)

    pid = _new_draft_project(hub)
    draft.put(id="dhub4", title="T", project=pid)
    draft.put(
        id="dhub4",
        chunk_kind="paragraph",
        text=f"Living citation [{finding_handle}].",
        at={"last": True},
    )
    ref = hub.live_store.get_ref(kind="draft", id="dhub4")
    out = tmp_path / "dhub4.docx"
    res = export_docx(hub.live_store, ref, target_path=out)
    assert res.cited_slugs == ["docxa21", "docxb22"]

    hrefs = {
        r.target_ref
        for r in docx.Document(str(out)).part.rels.values()
        if r.reltype.endswith("hyperlink")
    }
    assert "https://doi.org/10.1/docxa21" in hrefs
    assert "https://doi.org/10.1/docxb22" in hrefs
    assert sum("uol.primo.exlibrisgroup.com" in h for h in hrefs) == 2

    out_off = tmp_path / "dhub4-off.docx"
    export_docx(
        hub.live_store,
        ref,
        target_path=out_off,
        doi_links=False,
        library_links=False,
    )
    rels_off = docx.Document(str(out_off)).part.rels
    assert not any(r.reltype.endswith("hyperlink") for r in rels_off.values())


def test_doi_and_library_link_switches_are_independent(
    draft: DraftHandler, hub: Hub, tmp_path: Path
) -> None:
    """``doi_links`` and ``library_links`` are independent: doi-on/
    library-off emits only the doi hyperlink; doi-off/library-on emits
    only the library-search hyperlink."""
    a = hub.live_store.insert_ref(
        kind="paper", slug="docxindep1", title="A", year=2001, meta={}
    ).id
    hub.live_store.insert_ref_identifiers(a, [("doi", "10.1/docxindep1", "manual")])

    pid = _new_draft_project(hub)
    draft.put(id="dindep", title="T", project=pid)
    draft.put(
        id="dindep",
        chunk_kind="paragraph",
        text="See [§docxindep1].",
        at={"last": True},
    )
    ref = hub.live_store.get_ref(kind="draft", id="dindep")

    out_doi_only = tmp_path / "dindep-doi.docx"
    export_docx(
        hub.live_store,
        ref,
        target_path=out_doi_only,
        doi_links=True,
        library_links=False,
    )
    hrefs_doi_only = {
        r.target_ref
        for r in docx.Document(str(out_doi_only)).part.rels.values()
        if r.reltype.endswith("hyperlink")
    }
    assert "https://doi.org/10.1/docxindep1" in hrefs_doi_only
    assert not any("uol.primo.exlibrisgroup.com" in h for h in hrefs_doi_only)

    out_lib_only = tmp_path / "dindep-lib.docx"
    export_docx(
        hub.live_store,
        ref,
        target_path=out_lib_only,
        doi_links=False,
        library_links=True,
    )
    hrefs_lib_only = {
        r.target_ref
        for r in docx.Document(str(out_lib_only)).part.rels.values()
        if r.reltype.endswith("hyperlink")
    }
    assert "https://doi.org/10.1/docxindep1" not in hrefs_lib_only
    assert any("uol.primo.exlibrisgroup.com" in h for h in hrefs_lib_only)


def test_hub_finding_doi_hyperlink_percent_encodes_reserved_chars(
    draft: DraftHandler, hub: Hub, tmp_path: Path
) -> None:
    """A legacy DOI with reserved URL chars (``<`` / ``>`` / ``;`` here)
    must be percent-encoded in the hyperlink relationship target — a raw
    ``<`` is not well-formed URL syntax."""
    from precis.taproot.hub import attach_evidence
    from precis.utils import handle_registry

    hub_ref = _mint_hub_claim(hub.live_store)
    a = hub.live_store.insert_ref(
        kind="paper", slug="docxa41", title="A — legacy DOI", year=1998, meta={}
    ).id
    citer = hub.live_store.insert_ref(
        kind="paper", slug="docxc49", title="Citer", year=2009, meta={}
    ).id
    for p in (a, citer):
        attach_evidence(
            hub.live_store, hub_ref_id=hub_ref, paper_ref_id=p, role="corroborates"
        )
    hub.live_store.add_link(src_ref_id=citer, dst_ref_id=a, relation="cites")
    hub.live_store.insert_ref_identifiers(
        a,
        [
            (
                "doi",
                "10.1002/(SICI)1097-0258(19980815/30)17:15/16<1661::AID-SIM968>3.0.CO;2-2",
                "manual",
            )
        ],
    )
    finding_handle = handle_registry.format_handle("finding", hub_ref)

    pid = _new_draft_project(hub)
    draft.put(id="dhub5", title="T", project=pid)
    draft.put(
        id="dhub5",
        chunk_kind="paragraph",
        text=f"Living citation [{finding_handle}].",
        at={"last": True},
    )
    ref = hub.live_store.get_ref(kind="draft", id="dhub5")
    out = tmp_path / "dhub5.docx"
    res = export_docx(hub.live_store, ref, target_path=out)
    assert res.cited_slugs == ["docxa41"]

    hrefs = {
        r.target_ref
        for r in docx.Document(str(out)).part.rels.values()
        if r.reltype.endswith("hyperlink")
    }
    assert any("%3C1661" in h for h in hrefs)
    assert not any("<" in h or ">" in h or ";" in h for h in hrefs)


def test_hub_finding_no_evidence_renders_no_cite(
    draft: DraftHandler, hub: Hub, tmp_path: Path
) -> None:
    from precis.utils import handle_registry

    hub_ref = _mint_hub_claim(hub.live_store)
    finding_handle = handle_registry.format_handle("finding", hub_ref)

    pid = _new_draft_project(hub)
    draft.put(id="dhub3", title="T", project=pid)
    draft.put(
        id="dhub3",
        chunk_kind="paragraph",
        text=f"Pending [{finding_handle}] evidence.",
        at={"last": True},
    )
    ref = hub.live_store.get_ref(kind="draft", id="dhub3")
    out = tmp_path / "dhub3.docx"
    res = export_docx(hub.live_store, ref, target_path=out)

    assert res.cited_slugs == []
    text = "\n".join(p.text for p in docx.Document(str(out)).paragraphs)
    assert "[1]" not in text
    assert "References" not in text


# ── Taproot Phase 2 — authorial pins reach draft export (docx) ─────────
# Mirrors the LaTeX exporter's pin coverage (tests/test_export_latex.py) —
# the SAME shared `precis.taproot.cite.apply_pin` policy applied through
# the `mentions` grammar's optional `pin` capture group.


def test_pin_replace_renders_pinned_not_derived_originator(
    draft: DraftHandler, hub: Hub, tmp_path: Path
) -> None:
    from precis.taproot.hub import attach_evidence
    from precis.utils import handle_registry

    hub_ref = _mint_hub_claim(hub.live_store)
    origin = hub.live_store.insert_ref(
        kind="paper", slug="docxp01", title="Original report", year=2001, meta={}
    ).id
    follow = hub.live_store.insert_ref(
        kind="paper", slug="docxq01", title="Follow-up", year=2005, meta={}
    ).id
    attach_evidence(
        hub.live_store, hub_ref_id=hub_ref, paper_ref_id=origin, role="corroborates"
    )
    attach_evidence(
        hub.live_store, hub_ref_id=hub_ref, paper_ref_id=follow, role="corroborates"
    )
    hub.live_store.add_link(src_ref_id=follow, dst_ref_id=origin, relation="cites")
    pinned = hub.live_store.insert_ref(
        kind="paper", slug="docxr01", title="Author's pick", meta={}
    ).id
    finding_handle = handle_registry.format_handle("finding", hub_ref)
    pin_handle = handle_registry.format_handle("paper", pinned)

    pid = _new_draft_project(hub)
    draft.put(id="dpin1", title="T", project=pid)
    draft.put(
        id="dpin1",
        chunk_kind="paragraph",
        text=f"Living citation [{finding_handle}>{pin_handle}].",
        at={"last": True},
    )
    ref = hub.live_store.get_ref(kind="draft", id="dpin1")
    out = tmp_path / "dpin1.docx"
    res = export_docx(hub.live_store, ref, target_path=out)

    assert res.cited_slugs == ["docxr01"]  # pinned handle, not the derived originator
    assert any("reconsider" in w for w in res.warnings)
    text = "\n".join(p.text for p in docx.Document(str(out)).paragraphs)
    assert "Author's pick" in text
    assert "Original report" not in text


def test_pin_supplement_renders_derived_plus_pinned(
    draft: DraftHandler, hub: Hub, tmp_path: Path
) -> None:
    from precis.taproot.hub import attach_evidence
    from precis.utils import handle_registry

    hub_ref = _mint_hub_claim(hub.live_store)
    origin = hub.live_store.insert_ref(
        kind="paper", slug="docxp02", title="Original report", year=2001, meta={}
    ).id
    follow = hub.live_store.insert_ref(
        kind="paper", slug="docxq02", title="Follow-up", year=2005, meta={}
    ).id
    attach_evidence(
        hub.live_store, hub_ref_id=hub_ref, paper_ref_id=origin, role="corroborates"
    )
    attach_evidence(
        hub.live_store, hub_ref_id=hub_ref, paper_ref_id=follow, role="corroborates"
    )
    hub.live_store.add_link(src_ref_id=follow, dst_ref_id=origin, relation="cites")
    pinned = hub.live_store.insert_ref(
        kind="paper", slug="docxr02", title="Extra evidence", meta={}
    ).id
    finding_handle = handle_registry.format_handle("finding", hub_ref)
    pin_handle = handle_registry.format_handle("paper", pinned)

    pid = _new_draft_project(hub)
    draft.put(id="dpin2", title="T", project=pid)
    draft.put(
        id="dpin2",
        chunk_kind="paragraph",
        text=f"Living citation [{finding_handle}+{pin_handle}].",
        at={"last": True},
    )
    ref = hub.live_store.get_ref(kind="draft", id="dpin2")
    out = tmp_path / "dpin2.docx"
    res = export_docx(hub.live_store, ref, target_path=out)

    assert res.cited_slugs == ["docxp02", "docxr02"]  # derived + pinned, both present
    assert not any("reconsider" in w for w in res.warnings)
    text = "\n".join(p.text for p in docx.Document(str(out)).paragraphs)
    assert "[1]" in text and "[2]" in text


def test_pin_on_non_hub_finding_ignored_with_warning(
    draft: DraftHandler, hub: Hub, tmp_path: Path
) -> None:
    from precis.handlers.finding import FindingHandler

    _seed_paper(hub.live_store, "docxn01", "Established plain finding", 2019)
    resp = FindingHandler(hub=hub).put(
        title="t", body="b", scope={}, cited_in="docxn01"
    )
    ref_id = int(resp.body.split("id=")[1].split()[0].rstrip(",.()"))
    hub.live_store.update_ref(ref_id, meta_patch={"primary_cite_key": "docxn01"})
    from precis.utils import handle_registry

    finding_handle = handle_registry.format_handle("finding", ref_id)

    pid = _new_draft_project(hub)
    draft.put(id="dpin3", title="T", project=pid)
    draft.put(
        id="dpin3",
        chunk_kind="paragraph",
        text=f"Established [{finding_handle}>pa5].",
        at={"last": True},
    )
    ref = hub.live_store.get_ref(kind="draft", id="dpin3")
    out = tmp_path / "dpin3.docx"
    res = export_docx(hub.live_store, ref, target_path=out)

    assert res.cited_slugs == ["docxn01"]  # unchanged — pin is meaningless here
    assert any("pin on" in w and "ignored" in w for w in res.warnings)


def _si_docx_ctx(monkeypatch, *, with_parent: bool):
    import contextlib
    from types import SimpleNamespace

    from precis.export import docx as dx
    from precis.store import si_links

    parent = SimpleNamespace(id=1, slug="parent24", kind="paper", pdf_role=None)
    si = SimpleNamespace(id=2, slug="parent24si", kind="paper", pdf_role="supplement")
    store = _RefStore({("paper", "parent24"): parent, ("paper", "parent24si"): si})
    store.pool = type(
        "P", (), {"connection": lambda s: contextlib.nullcontext(object())}
    )()
    monkeypatch.setattr(
        si_links,
        "supplement_parent",
        lambda _conn, rid: (1, "parent24") if with_parent and rid == 2 else None,
    )
    return dx._Ctx(
        store=store,
        known_handles=set(),
        doi_links=False,
        library_links=False,
        library_label="x",
        library_search_url="x",
    )


def test_si_cite_numbers_as_parent_with_si_note(monkeypatch) -> None:
    import docx

    from precis.export import docx as dx

    ctx = _si_docx_ctx(monkeypatch, with_parent=True)
    para = docx.Document().add_paragraph()
    dx._cite("parent24si", ctx, para)
    assert ctx.cited == ["parent24"]
    assert para.text == "[1] (SI)"


def test_si_cite_without_parent_falls_back_with_warning_docx(monkeypatch) -> None:
    import docx

    from precis.export import docx as dx

    ctx = _si_docx_ctx(monkeypatch, with_parent=False)
    para = docx.Document().add_paragraph()
    dx._cite("parent24si", ctx, para)
    assert ctx.cited == ["parent24si"]
    assert para.text == "[1]"
    assert any("no live parent" in w for w in ctx.warnings)


def test_si_mark_survives_after_plain_mark_of_same_parent(monkeypatch) -> None:
    import docx

    from precis.export import docx as dx

    ctx = _si_docx_ctx(monkeypatch, with_parent=True)
    para = docx.Document().add_paragraph()
    dx._cite("parent24", ctx, para)
    dx._cite("parent24si", ctx, para)  # same parent, SI flag differs
    dx._cite("parent24si", ctx, para)  # identical consecutive -> collapses
    assert para.text == "[1][1] (SI)"


def test_mathrm_renders_upright_omml() -> None:
    """``\\mathrm{C}_{60}`` (the repaired ``C$_{60}$``) must come out as
    upright (``m:sty p``) OMML runs, single- and multi-character alike."""
    pytest.importorskip("latex2mathml")

    from precis.export.omml import latex_to_omml

    ns = {"m": "http://schemas.openxmlformats.org/officeDocument/2006/math"}
    for src, base in ((r"\mathrm{C}_{60}", "C"), (r"\mathrm{WS}_2", "WS")):
        omath = latex_to_omml(src)
        assert omath is not None
        runs = omath.xpath("//m:e/m:r", namespaces=ns)
        assert "".join(r.xpath("string(m:t)", namespaces=ns) for r in runs) == base
        assert all(r.xpath("m:rPr/m:sty/@m:val", namespaces=ns) == ["p"] for r in runs)


# ── fidelity vs. the LaTeX exporter (docx-fidelity pass) ───────────────


class _NonHubConn:
    def execute(self, *_a, **_k):
        return self

    def fetchone(self):
        return None


class _NonHubPool:
    def connection(self):
        import contextlib

        return contextlib.nullcontext(_NonHubConn())


class _FindingFake:
    """Store for a plain finding handle with no cite key of its own."""

    def __init__(self, refs):
        self._refs = refs
        self.pool = _NonHubPool()

    def fetch_refs_by_ids(self, ids):
        from types import SimpleNamespace

        return {i: self._refs.get(i, SimpleNamespace(meta={})) for i in ids}

    def tags_for(self, _ref_id):
        return []

    def get_ref(self, **_k):
        return None


def _render_text(text: str, ctx) -> str:
    from precis.export import docx as dx

    para = docx.Document().add_paragraph()
    dx._render_inline(text, ctx, para)
    return para.text


def test_conjunction_hub_cites_union_of_conjunct_sources(monkeypatch) -> None:
    """A finding with no cite key of its own (conjunction hub) cites its
    atoms' sources, like LaTeX — it used to vanish from docx."""
    from precis.export import docx as dx

    monkeypatch.setattr(dx, "_conjunct_cite_keys", lambda _s, _pk: ["a21", "b22"])
    monkeypatch.setattr(dx, "_render_trust_mark", lambda *_a: None)
    ctx = dx._Ctx(
        store=_FindingFake({}),
        known_handles=set(),
        doi_links=False,
        library_links=False,
        library_label="x",
        library_search_url="x",
    )
    out = _render_text("claim [fi7].", ctx)
    assert out == "claim [1][2]."
    assert ctx.cited == ["a21", "b22"]


def test_sourceless_finding_cite_pulls_space_before_punctuation(monkeypatch) -> None:
    from precis.export import docx as dx

    monkeypatch.setattr(dx, "_conjunct_cite_keys", lambda _s, _pk: [])
    monkeypatch.setattr(dx, "_render_trust_mark", lambda *_a: None)
    ctx = dx._Ctx(
        store=_FindingFake({}),
        known_handles=set(),
        doi_links=False,
        library_links=False,
        library_label="x",
        library_search_url="x",
    )
    out = _render_text("in frameworks [fi7]. And (~130 GPa [fi7]) end", ctx)
    assert out == "in frameworks. And (~130 GPa) end"
    assert any("no citable source" in w for w in ctx.warnings)


def test_cite_run_does_not_collapse_across_paragraphs(monkeypatch) -> None:
    """A paragraph that opens with the paper the previous one closed on keeps
    its mark: consecutive-cite collapse is scoped to one chunk."""
    from precis.export import docx as dx

    ctx = dx._Ctx(
        store=None,
        known_handles=set(),
        doi_links=False,
        library_links=False,
        library_label="x",
        library_search_url="x",
    )
    assert _render_text("one [§smith20].", ctx) == "one [1]."
    assert _render_text("[§smith20] two", ctx) == "[1] two"


def test_si_cite_via_alias_shares_the_parents_number(monkeypatch) -> None:
    """The same paper cited plainly (under an alias key) and via its SI
    record must get ONE number; SI keeps the '(SI)' postnote."""
    from types import SimpleNamespace

    import docx as _docx

    from precis.export import docx as dx

    ctx = _si_docx_ctx(monkeypatch, with_parent=True)
    # an alias key resolves to the same ref whose canonical slug is parent24
    ctx.store._refs[("paper", "parent24alias")] = SimpleNamespace(
        id=1, slug="parent24", kind="paper", pdf_role=None
    )
    para = _docx.Document().add_paragraph()
    dx._cite("parent24alias", ctx, para)
    para.add_run(" and ")
    dx._cite("parent24si", ctx, para)
    assert ctx.cited == ["parent24"]
    assert para.text == "[1] and [1] (SI)"


def test_xref_numbers_follow_latex_section_counting() -> None:
    """The seeded title is skipped and the nested body lifted a level
    (Intro is 1, a depth-2 child 1.1); the draft's own Glossary heading is
    skipped; figures count in reading order; run-in headings (depth >= 3)
    resolve to their enclosing section."""
    from types import SimpleNamespace as NS

    from precis.export.docx import _xref_numbers

    def ch(dc, kind, depth, text=""):
        return NS(dc=dc, chunk_kind=kind, depth=depth, text=text)

    chunks = [
        ch("dc1", "heading", 0, "Title"),
        ch("dc2", "heading", 1, "Intro"),
        ch("dc3", "figure", 2),
        ch("dc4", "heading", 1, "Applications"),
        ch("dc5", "heading", 2, "Catalysis"),
        ch("dc6", "heading", 2, "Batteries"),
        ch("dc7", "heading", 3, "Run-in"),
        ch("dc8", "figure", 2),
        ch("dc9", "heading", 0, "Glossary"),
        ch("dc10", "term", 1, "x"),
        ch("dc11", "heading", 4, "Run-in"),
    ]
    xref, shown = _xref_numbers(chunks, "  title ")
    assert "dc1" not in xref  # the seeded title heading is not a section
    assert xref["dc2"] == ("section", "1")
    assert xref["dc4"] == ("section", "2")
    assert xref["dc6"] == ("section", "2.2")
    assert xref["dc7"] == ("section", "2.2.1")
    # sections beside the title at depth 0 (no lift): nothing is shifted
    flat = [ch("dc1", "heading", 0, "Title"), ch("dc2", "heading", 0, "Intro")]
    assert _xref_numbers(flat, "Title")[0]["dc2"] == ("section", "1")
    # an untitled lookup keeps the old behaviour: depth-0 heading counts
    assert _xref_numbers(chunks)[0]["dc1"] == ("section", "1")
    assert xref["dc3"] == ("fig.", "1") and xref["dc8"] == ("fig.", "2")
    assert "dc9" not in xref and "dc11" not in shown and "dc7" in shown


def test_xref_numbers_resolve_paragraphs_and_back_matter_like_latex() -> None:
    """A cross-ref to a plain paragraph prints its enclosing section (what
    ``\\cref`` does for a label inside a section); journal back matter is
    unnumbered and a cross-ref to it prints its name (``\\nameref``)."""
    from types import SimpleNamespace as NS

    from precis.export.docx import _xref_numbers

    def ch(dc, kind, depth, text=""):
        return NS(dc=dc, chunk_kind=kind, depth=depth, text=text)

    chunks = [
        ch("abs", "paragraph", 0, "The abstract."),
        ch("h1", "heading", 0, "Introduction"),
        ch("p1", "paragraph", 1, "body"),
        ch("h11", "heading", 1, "Scope"),
        ch("p2", "paragraph", 2, "body"),
        ch("h2", "heading", 0, "Results"),
        ch("ul", "ulist", 1),
        ch("it", "item", 2, "an item"),
        ch("bm", "heading", 0, "Author Information"),
        ch("bm1", "heading", 1, "Notes"),
        ch("p3", "paragraph", 1, "thanks"),
    ]
    xref, shown = _xref_numbers(chunks, "Paper")
    assert "abs" not in xref  # front matter: nothing to print
    assert xref["p1"] == ("section", "1")
    assert xref["p2"] == ("section", "1.1")
    assert xref["it"] == ("section", "2") and "ul" not in xref
    assert xref["bm"] == ("", "Author Information")
    assert xref["bm1"] == ("", "Notes")
    assert "bm" not in shown and "bm1" not in shown
    assert xref["p3"] == ("section", "2")  # \section* does not step the counter


def test_docx_without_seeded_title_chunk_prints_ref_title_first(
    draft: DraftHandler, hub: Hub, tmp_path: Path
) -> None:
    """A draft whose reading order opens with the abstract (the seeded
    title heading retired) still gets the title block up front, and its
    first depth-0 section is numbered "1", not styled as the title."""
    pid = _new_draft_project(hub)
    draft.put(id="dnt", title="Nanobuds", project=pid)
    ref = hub.live_store.get_ref(kind="draft", id="dnt")
    assert ref is not None
    title_h = hub.live_store.drafts.reading_order(ref.id)[0].handle
    draft.put(
        id="dnt", chunk_kind="paragraph", text="The abstract.", at={"first": True}
    )
    draft.put(id="dnt", chunk_kind="heading", text="Introduction", at={"last": True})
    draft.put(id="dnt", chunk_kind="paragraph", text="Body.", at={"last": True})
    draft.delete(id=f"¶{title_h}")
    out = tmp_path / "dnt.docx"
    export_docx(hub.live_store, ref, target_path=out)
    paras = docx.Document(str(out)).paragraphs
    texts = [p.text for p in paras]
    assert texts[0] == "Nanobuds" and paras[0].style.name == "Title"
    assert texts.index("Abstract") < texts.index("The abstract.")
    intro = next(p for p in paras if p.text.endswith("Introduction"))
    assert intro.text == "1 Introduction" and intro.style.name == "Heading 1"


def test_bare_dc_crossref_prints_section_and_figure_numbers() -> None:
    from precis.export import docx as dx

    ctx = dx._Ctx(
        store=None,
        known_handles=set(),
        xref={"dc4": ("section", "1.5.5"), "dc3": ("fig.", "1")},
        dc_handles={"dc3", "dc4", "dc5"},
        legacy_to_dc={"Abc": "dc4"},
    )
    assert (
        _render_text("(see [dc4]) and [¶dc4]", ctx)
        == "(see section 1.5.5) and section 1.5.5"
    )
    assert _render_text("Shown in [dc3].", ctx) == "Shown in fig. 1."
    # sentence start capitalises
    assert (
        _render_text("[dc4] covers it. [dc3] too", ctx)
        == "Section 1.5.5 covers it. Fig. 1 too"
    )
    assert _render_text("via [¶Abc]", ctx) == "via section 1.5.5"
    # authored surface wins; a non-live chunk downgrades with a warning
    assert _render_text("[the intro](dc4)", ctx) == "the intro"
    assert _render_text("see [dc99]", ctx) == "see dc99"
    assert any("dc99" in w for w in ctx.warnings)


def test_export_numbers_headings_and_resolves_crossrefs(
    draft: DraftHandler, hub: Hub, tmp_path: Path
) -> None:
    import re

    pid = _new_draft_project(hub)
    draft.put(id="dxr", title="T", project=pid)
    h1 = draft.put(id="dxr", chunk_kind="heading", text="Intro", at={"last": True})
    h1_match = re.search(r"dc\d+", h1.body)
    assert h1_match is not None
    h1_dc = h1_match.group(0)
    draft.put(
        id="dxr",
        chunk_kind="paragraph",
        text=f"As argued in [{h1_dc}], things hold (see [¶{h1_dc}]).",
        at={"last": True},
    )
    ref = hub.live_store.get_ref(kind="draft", id="dxr")
    out = tmp_path / "dxr.docx"
    export_docx(hub.live_store, ref, target_path=out)
    text = "\n".join(p.text for p in docx.Document(str(out)).paragraphs)
    assert "As argued in section" in text
    assert "(see section" in text
    assert "Intro" in text


def test_aligned_equation_converts_one_omml_line_per_row() -> None:
    from precis.export.docx import _render_equation

    doc = docx.Document()
    body = (
        r" \begin{aligned} \sum_{n=4}^{8} (6 - n)\,P_n &= 6\chi, \\ "
        r"2P_4 + P_5 - P_7 - 2P_8 &= 6\chi. \end{aligned} "
    )
    _render_equation(doc, body, 2, False)
    paras = doc.paragraphs
    assert len(paras) == 2
    xml = "".join(p._p.xml for p in paras)
    assert "begin{aligned}" not in xml and "&amp;" not in xml
    assert xml.count("<m:oMath") == 2  # both rows went through OMML
    assert "(2)" not in paras[0].text and paras[1].text.endswith("(2)")
    # unnumbered (starred) variant: rows, no label
    doc2 = docx.Document()
    _render_equation(doc2, body, None, True)
    assert len(doc2.paragraphs) == 2
    assert not any(" + (" in p.text for p in doc2.paragraphs)


def test_single_line_equation_unchanged_by_aligned_support() -> None:
    from precis.export.docx import _render_equation

    doc = docx.Document()
    _render_equation(doc, r" \chi = V - E + F. ", 1, False)
    assert len(doc.paragraphs) == 1 and doc.paragraphs[0].text.endswith("(1)")


def _abbrev_ctx(abbrevs):
    from precis.export import docx as dx

    return dx._Ctx(store=None, known_handles=set(), abbrevs=abbrevs)


def test_abbrev_plural_first_use_pluralises_the_long_form() -> None:
    ctx = _abbrev_ctx({"CNT": "carbon nanotube"})
    out = _render_text("We study CNTs. Later CNTs again.", ctx)
    assert out == "We study carbon nanotubes (CNTs). Later CNTs again."


def test_abbrev_spelled_out_in_prose_collapses_to_one_expansion() -> None:
    ctx = _abbrev_ctx(
        {
            "HA": "hexylamine",
            "pG": "pristine graphene",
            "CNT": "carbon nanotube",
        }
    )
    out = _render_text(
        "Using hexylamine (HA) and Pristine Graphene (pG) on carbon nanotubes (CNTs);"
        " HA again.",
        ctx,
    )
    assert out == (
        "Using hexylamine (HA) and pristine graphene (pG) on carbon nanotubes"
        " (CNTs); HA again."
    )
    assert "hexylamine (hexylamine" not in out


def test_abbrev_first_use_at_sentence_start_keeps_capital() -> None:
    ctx = _abbrev_ctx({"GGA": "generalized gradient approximation"})
    assert (
        _render_text("GGA works. The GGA fails.", ctx)
        == "Generalized gradient approximation (GGA) works. The GGA fails."
    )
    ctx2 = _abbrev_ctx({"GGA": "generalized gradient approximation"})
    assert (
        _render_text("It is cheap. GGA works.", ctx2)
        == "It is cheap. Generalized gradient approximation (GGA) works."
    )


def test_abbrev_followed_by_subscripted_index_keeps_the_subscript() -> None:
    """``NICS(1)$_{zz}$``: the shared empty-base repair pulled ``1)`` into the
    math, splitting the parenthesis; a pure-script span is a sub run now."""
    from precis.export import docx as dx

    ctx = _abbrev_ctx({"NICS": "nucleus-independent chemical shift"})
    para = docx.Document().add_paragraph()
    dx._render_inline("a NICS(1)$_{zz}$ index and H$_2$O and x$^{2+}$", ctx, para)
    assert para.text == (
        "a nucleus-independent chemical shift (NICS)(1)zz index and H2O and x2+"
    )
    subs = [r.text for r in para.runs if r.font.subscript]
    sups = [r.text for r in para.runs if r.font.superscript]
    assert subs == ["zz", "2"] and sups == ["2+"]
    assert not para._p.xpath(".//m:oMath")


def test_figure_captions_carry_figure_n_label(
    draft: DraftHandler, hub: Hub, tmp_path: Path
) -> None:
    """Captions open with a bold "Figure N:" in reading order, and a bare
    [dc<id>] cross-ref to the second figure prints the same number."""
    pid = _new_draft_project(hub)
    draft.put(id="dfig", title="T", project=pid)
    ref = hub.live_store.get_ref(kind="draft", id="dfig")
    assert ref is not None
    title_h = hub.live_store.drafts.reading_order(ref.id)[0].handle
    f1 = hub.live_store.drafts.add_chunks(
        ref_id=ref.id,
        chunk_kind="figure",
        text="First widget.",
        at={"after": title_h},
        split=False,
    )[0]
    f2 = hub.live_store.drafts.add_chunks(
        ref_id=ref.id,
        chunk_kind="figure",
        text="Second widget.",
        at={"after": f1.handle},
        split=False,
    )[0]
    dc2 = hub.live_store.drafts.reading_order(ref.id)[2].dc
    assert f2.chunk_id
    draft.put(
        id="dfig",
        chunk_kind="paragraph",
        text=f"As in [{dc2}].",
        at={"last": True},
    )
    out = tmp_path / "dfig.docx"
    export_docx(hub.live_store, ref, target_path=out)
    paras = docx.Document(str(out)).paragraphs
    texts = [p.text for p in paras]
    assert "Figure 1: First widget." in texts
    assert "Figure 2: Second widget." in texts
    assert "As in fig. 2." in texts
    cap = next(p for p in paras if p.text.startswith("Figure 1:"))
    assert cap.runs[0].bold and cap.runs[1].italic and not cap.runs[0].italic


def test_bib_markup_runs_convert_html_and_math_to_styles() -> None:
    from precis.export.docx import _bib_markup_runs

    runs = _bib_markup_runs(
        "{C$_{60}$} and C<sub>60</sub>: <i>In situ</i> <scp>iii</scp> Fe$^{3+}$ &amp; x"
    )
    text = "".join(t for t, _ in runs)
    assert text == "C60 and C60: In situ iii Fe3+ & x"
    styled = {t: st for t, st in runs if st}
    assert styled["60"] == frozenset({"sub"})
    assert styled["In situ"] == frozenset({"i"})
    assert styled["iii"] == frozenset({"sc"})
    assert styled["3+"] == frozenset({"sup"})
    for leak in ("$", "{", "<", "&amp;"):
        assert leak not in text


def test_reference_line_carries_venue_fields_and_no_double_period() -> None:
    from types import SimpleNamespace

    from precis.export.docx import _format_reference, _reference_runs

    paper = SimpleNamespace(
        id=3,
        slug="doe20",
        kind="paper",
        title="Charge transfer in {C$_{60}$} fields.",
        authors=[{"name": "Doe, Jane"}, {"name": "Roe & Sons"}],
        year=2020,
        meta={
            "venue": "J. Phys. <scp>iii</scp>",
            "volume": "12",
            "number": "4",
            "pages": "100-110",
        },
    )
    store = _RefStore({("paper", "doe20"): paper})
    warnings: list[str] = []
    line = _format_reference(store, "doe20", warnings)
    assert line == (
        "Doe, Jane; Roe & Sons (2020). Charge transfer in C60 fields."
        " J. Phys. iii 12(4), 100-110."
    )
    assert ".." not in line
    runs = _reference_runs(store, "doe20", warnings)
    assert ("60", frozenset({"sub"})) in runs
    assert ("J. Phys. ", frozenset({"i"})) in runs
    assert ("iii", frozenset({"i", "sc"})) in runs


def test_references_section_renders_sub_and_italic_runs(tmp_path: Path) -> None:
    from types import SimpleNamespace

    from precis.export import docx as dx

    paper = SimpleNamespace(
        id=3,
        slug="doe20",
        kind="paper",
        title="On C<sub>60</sub>",
        authors=[{"name": "Doe, Jane"}],
        year=2020,
        meta={"journal": "Carbon", "volume": "7"},
    )
    ctx = dx._Ctx(
        store=_RefStore({("paper", "doe20"): paper}),
        known_handles=set(),
        cited=["doe20"],
    )
    doc = docx.Document()
    dx._append_references(doc, ctx)
    p = doc.paragraphs[-1]
    assert p.text == "[1] Doe, Jane (2020). On C60. Carbon 7."
    assert any(r.font.subscript and r.text == "60" for r in p.runs)
    assert any(r.italic and r.text == "Carbon" for r in p.runs)


def test_title_page_has_date_line_and_abstract_heading(
    draft: DraftHandler, hub: Hub, tmp_path: Path
) -> None:
    from datetime import UTC, datetime

    pid = _new_draft_project(hub)
    draft.put(id="dab", title="Paper", project=pid)
    draft.put(id="dab", chunk_kind="paragraph", text="The abstract.", at={"last": True})
    draft.put(id="dab", chunk_kind="heading", text="Intro", at={"last": True})
    ref = hub.live_store.get_ref(kind="draft", id="dab")
    out = tmp_path / "dab.docx"
    export_docx(hub.live_store, ref, target_path=out)
    paras = docx.Document(str(out)).paragraphs
    texts = [p.text for p in paras]
    today = datetime.now(UTC)
    date_line = f"{today:%B} {today.day}, {today.year}"
    assert texts[1] == date_line  # right under the title (no byline here)
    i = texts.index("Abstract")
    assert texts[i + 1] == "The abstract."
    assert texts[0] == "Paper" and i == 2
    assert next(p for p in paras if p.text == "Abstract").style.name.startswith(
        "Heading"
    )


def test_nested_layout_numbers_sections_from_one(
    draft: DraftHandler, hub: Hub, tmp_path: Path
) -> None:
    """title -> Introduction -> Scope displays "1" and "1.1", and a
    cross-ref to Scope prints the same number."""
    import re

    pid = _new_draft_project(hub)
    draft.put(id="dnest", title="Nested", project=pid)
    ref = hub.live_store.get_ref(kind="draft", id="dnest")
    assert ref is not None
    title_dc = hub.live_store.drafts.reading_order(ref.id)[0].handle
    intro = draft.put(
        id="dnest",
        chunk_kind="heading",
        text="Introduction",
        at={"into": title_dc, "last": True},
    )
    intro_m = re.search(r"dc\d+", intro.body)
    assert intro_m is not None
    intro_h = intro_m.group(0)
    scope = draft.put(
        id="dnest",
        chunk_kind="heading",
        text="Scope",
        at={"into": intro_h, "last": True},
    )
    scope_m = re.search(r"dc\d+", scope.body)
    assert scope_m is not None
    scope_h = scope_m.group(0)
    draft.put(
        id="dnest",
        chunk_kind="paragraph",
        text=f"See [{scope_h}].",
        at={"into": scope_h, "last": True},
    )
    out = tmp_path / "dnest.docx"
    export_docx(hub.live_store, ref, target_path=out)
    paras = docx.Document(str(out)).paragraphs
    texts = [p.text for p in paras]
    assert "Nested" in texts and "1 Introduction" in texts and "1.1 Scope" in texts
    assert "See section 1.1." in texts
    head = next(p for p in paras if p.text == "1.1 Scope")
    assert head.style.name == "Heading 2"
