"""gr491599: one first-use abbreviation expander for docx and LaTeX.

Each case renders the same source through both exporters' real inline
paths and pins both outputs, so a rule can no longer hold in one format
and not the other.
"""

from __future__ import annotations

import docx
import pytest

from precis.export import _abbrev
from precis.export import docx as dx
from precis.export import latex as lx

AB = {
    "BN": "boron nitride",
    "CNB": "carbon nanobud",
    "DFT": "density functional theory",
    "DOS": "density of states",
    "FCNB": "functionalized CNB",
    "GNR": "graphene nanoribbon",
    "GNR-FET": "graphene nanoribbon field-effect transistor",
    "HOMO": "highest occupied molecular orbital",
    "LUMO": "lowest unoccupied molecular orbital",
    "PBE": "Perdew-Burke-Ernzerhof",
    "XPS": "X-ray photoelectron spectroscopy",
}


def _docx_ctx() -> dx._Ctx:
    return dx._Ctx(
        store=None,
        known_handles=set(),
        abbrevs=dict(AB),
        library_label="x",
        library_search_url="x",
    )


def _tex_ctx() -> lx._Ctx:
    return lx._Ctx(
        keymap=lx._acronym_keymap(AB),
        abbrevs=dict(AB),
        known_handles=set(),
        store=None,
        library_label="x",
        library_search_url="x",
    )


def _both(*texts: str) -> tuple[list[str], list[str], dx._Ctx]:
    dctx, tctx = _docx_ctx(), _tex_ctx()
    dout, tout = [], []
    for text in texts:
        p = docx.Document().add_paragraph()
        dx._render_inline(text, dctx, p)
        dout.append(p.text)
        tout.append(lx._render_inline(text, tctx))
    return dout, tout, dctx


@pytest.mark.parametrize(
    ("src", "want_docx", "want_tex"),
    [
        # M7: "et al." / "e.g." / "Fig." do not end the sentence
        (
            "Wu et al. DFT calculations show",
            "Wu et al. density functional theory (DFT) calculations show",
            r"Wu et al. \gls{dft} calculations show",
        ),
        (
            "solids, e.g. DFT runs",
            "solids, e.g. density functional theory (DFT) runs",
            r"solids, e.g. \gls{dft} runs",
        ),
        (
            "It fails. DFT works",
            "It fails. Density functional theory (DFT) works",
            r"It fails. \Gls{dft} works",
        ),
        # M6: nothing inside a URL or DOI expands or uses up the first use
        (
            "see https://example.org/DFT/data and 10.1000/DFT.7 for DFT inputs",
            "see https://example.org/DFT/data and 10.1000/DFT.7 for"
            " density functional theory (DFT) inputs",
            r"see https://example.org/DFT/data and 10.1000/DFT.7 for \gls{dft} inputs",
        ),
        # L1: a hyphen compound is a use; a hyphen prefix is not
        (
            "CNB-based films on h-BN",
            "Carbon nanobud (CNB)-based films on h-BN",
            r"\Gls{cnb}-based films on h-BN",
        ),
        # a registered compound wins over its parts
        (
            "a GNR-FET",
            "a graphene nanoribbon field-effect transistor (GNR-FET)",
            r"a \gls{gnrfet}",
        ),
        # H3: a long form containing another short collapses once
        (
            "a functionalized CNB (FCNB) composite",
            "a functionalized CNB (FCNB) composite",
            r"a \gls{fcnb} composite",
        ),
        # nit: two new shorts joined by a dash are defined together
        (
            "the HOMO–LUMO gap",
            "the highest occupied molecular orbital–lowest unoccupied molecular"
            " orbital (HOMO–LUMO) gap",
            r"the \glsxtrlong{homo}{\textendash}\glsxtrlong{lumo}"
            r" (\glsxtrshort{homo}{\textendash}\glsxtrshort{lumo})"
            r"\glsunset{homo}\glsunset{lumo} gap",
        ),
        # nit: a first use inside parentheses brackets the short
        (
            "the choice (DFT, PBE) matters",
            "the choice (density functional theory [DFT],"
            " Perdew-Burke-Ernzerhof [PBE]) matters",
            r"the choice (\glsbrk{\gls}{dft}, \glsbrk{\gls}{pbe}) matters",
        ),
        # nit: an irregular plural
        (
            "XPSs and DOSs",
            "X-ray photoelectron spectroscopies (XPSs) and density of states (DOSs)",
            r"\Glspl{xps} and \glspl{dos}",
        ),
        # nit: a long form in link text defines the short that follows it
        (
            "[Density functional theory](https://x.org) (DFT) is used",
            "Density functional theory (DFT) is used",
            r"\href{https://x.org}{Density functional theory}"
            r" (\glsunset{dft}\gls{dft}) is used",
        ),
        # L3 parity: a definition opening the chunk in bold is capitalised
        (
            "**DFT** first",
            "Density functional theory (DFT) first",
            r"\textbf{\Gls{dft}} first",
        ),
    ],
)
def test_docx_and_latex_define_the_same_occurrence(
    src: str, want_docx: str, want_tex: str
) -> None:
    dout, tout, _ = _both(src)
    assert dout == [want_docx]
    assert tout == [want_tex]


def test_hyphen_compound_counts_as_the_use_and_lists_the_short() -> None:
    """L1: a short used only in compounds is defined once and lands in the
    docx acronym list; the later bare use stays short."""
    dout, tout, dctx = _both("CNB-based films grow.", "A CNB film.")
    assert dout == ["Carbon nanobud (CNB)-based films grow.", "A CNB film."]
    assert tout[1] == r"A \glstip{cnb} film."
    assert "CNB" in dctx.used_acr


def test_heading_prints_the_bare_short_and_keeps_the_first_use() -> None:
    """M2: a heading never defines; the body's first use still does."""
    dctx, tctx = _docx_ctx(), _tex_ctx()
    head = docx.Document().add_paragraph()
    dctx.in_heading = True
    dx._render_inline("DFT Benchmarks", dctx, head)
    dctx.in_heading = False
    body = docx.Document().add_paragraph()
    dx._render_inline("DFT is cheap.", dctx, body)
    assert head.text == "DFT Benchmarks"
    assert body.text == "Density functional theory (DFT) is cheap."
    assert "DFT" in dctx.used_acr

    tex_head = lx._render_moving_arg("DFT Benchmarks", tctx, heading=True)
    assert tex_head == r"\glsfmtshort{dft} Benchmarks"
    assert not tctx.in_heading
    assert lx._render_inline("DFT is cheap.", tctx) == r"\Gls{dft} is cheap."


def test_latex_heading_chunk_renders_bare_short() -> None:
    """M2 through the chunk walk: ``\\section{\\glsfmtshort{…}}``, and the
    paragraph after it carries the definition."""
    from types import SimpleNamespace

    ctx = _tex_ctx()
    chunks = [
        SimpleNamespace(
            chunk_kind="heading", text="DFT Benchmarks", depth=1, dc="dc1", meta={}
        ),
        SimpleNamespace(
            chunk_kind="paragraph", text="DFT is cheap.", depth=2, dc="dc2", meta={}
        ),
    ]
    tex = "\n".join(
        lx._render_chunk_lines(
            chunks,
            ctx,
            depth_of=lambda c: c.depth,
            title_dc=None,
            back_matter=frozenset(),
        )
    )
    # bare in the heading, but indexed after it so the acronym list has it
    assert r"section{\glsfmtshort{dft} Benchmarks}" in tex
    assert tex.count(r"\glsadd{dft}") == 1
    assert r"\Gls{dft} is cheap." in tex


def test_build_acronyms_sets_irregular_longplural() -> None:
    tex = lx.build_acronyms(
        {"XPS": "X-ray photoelectron spectroscopy", "MOF": "metal-organic framework"}
    )
    assert (
        r"\newacronym[longplural={X-ray photoelectron spectroscopies}]{xps}{XPS}" in tex
    )
    assert r"\newacronym{mof}{MOF}{metal-organic framework}" in tex


@pytest.mark.parametrize(
    ("before", "start"),
    [
        ("", True),
        ("It ends. ", True),
        ('He said "done." ', True),
        ("Wu et al. ", False),
        ("e.g. ", False),
        ("see Fig. ", False),
        ("as in Fig. 3. ", True),
        ("mid sentence ", False),
        ("It ends. **", True),
    ],
)
def test_sentence_start(before: str, start: bool) -> None:
    assert _abbrev.at_sentence_start(before) is start


def test_plural_long() -> None:
    assert _abbrev.plural_long("X-ray photoelectron spectroscopy") == (
        "X-ray photoelectron spectroscopies"
    )
    assert _abbrev.plural_long("density of states") == "density of states"
    assert _abbrev.plural_long("carbon nanotube") == "carbon nanotubes"
    assert _abbrev.plural_long("relay") == "relays"
