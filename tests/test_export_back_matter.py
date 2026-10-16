"""Back-matter order (gr493199) and the published-claims row layout
(gr493200): DB-free, fake stores. The DB-backed end-to-end order check
lives in ``test_export_docx.py``."""

from __future__ import annotations

import types
from datetime import UTC, datetime
from typing import Any

import docx as docx_lib

from precis.export import _back_matter, latex
from precis.export import docx as dx
from precis.export._nanopub_appendix import SECTION_TITLE
from tests.test_export_docx import _FindingFake
from tests.test_export_latex import _FindingStore

_ASM = dict(
    title="T", author_block="", body="body", acronyms="", appendix="", doc_type=""
)


def _tex(**kw: Any) -> str:
    return latex.assemble_document(**{**_ASM, **kw})


def test_latex_bibliography_before_glossaries() -> None:
    out = _tex()
    assert out.index("\\printbibliography") < out.index("\\printglossaries")


def test_latex_full_order_follows_constant() -> None:
    out = _tex(
        unverified_section="%UNV",
        data_package_section="%DPK",
        published_section="%PUB",
        appendix="%SRC",
    )
    marks = {
        "unverified": "%UNV",
        "data_package": "%DPK",
        "references": "\\printbibliography",
        "acronyms": "\\printglossaries",
        "published": "%PUB",
        "sources": "%SRC",
    }
    positions = [out.index(marks[n]) for n in _back_matter.BACK_MATTER_ORDER]
    assert positions == sorted(positions)


def test_latex_patent_mode_has_no_bibliography() -> None:
    out = _tex(doc_type="patent")
    assert "\\printbibliography" not in out
    assert "\\printglossaries" in out


def test_latex_follows_the_shared_constant(monkeypatch) -> None:
    """Swapping the constant swaps the output: latex keeps no order of its own.
    (The docx twin is in test_export_docx.py.)"""
    monkeypatch.setattr(
        _back_matter,
        "BACK_MATTER_ORDER",
        (
            "acronyms",
            "references",
            "unverified",
            "data_package",
            "published",
            "sources",
        ),
    )
    out = _tex()
    assert out.index("\\printglossaries") < out.index("\\printbibliography")


# ── published claims rows ──────────────────────────────────────────────


class _Row:
    state = "published"
    trusty_uri = "https://w3id.org/np/RAabc"
    approved_title = "Nanobuds hold."
    published_at = datetime(2026, 1, 2, tzinfo=UTC)
    artifact_id = None


class _NpStore:
    def nanopub_publish_row(self, _fid: int) -> _Row:
        return _Row()


def _trust(fids: list[int]) -> Any:
    return types.SimpleNamespace(cited=fids)


def test_published_entry_carries_fid() -> None:
    from precis.export._nanopub_appendix import published_claim_entries

    (e,) = published_claim_entries(_NpStore(), _trust([7]))
    assert e.fid == 7


def test_latex_bullet_cites_only_body_cited_keys(monkeypatch) -> None:
    monkeypatch.setattr(latex, "_trust_mark_latex", lambda _c, _r: "")
    monkeypatch.setattr(latex, "_conjunct_cite_keys", lambda _s, _pk: ["a21", "b22"])
    store = _FindingStore({7: types.SimpleNamespace(meta={})})
    ctx = latex._Ctx(keymap={}, known_handles=set(), store=store)
    latex._render_inline("claim [fi7].", ctx)
    assert ctx.claim_cites == {7: ["a21", "b22"]}
    cited_before = list(ctx.cited)

    # a key the body never cited must not appear even if offered
    out = latex.build_published_claims_section(
        _NpStore(), _trust([7]), {7: ctx.claim_cites[7]}
    )
    assert "\\item ``Nanobuds hold.'' ; " not in out
    assert "\\item ``Nanobuds hold.''; Sources: \\cite{a21,b22};" in out
    assert "\\url{https://w3id.org/np/RAabc} --- published 2026-01-02" in out
    assert out.count("\\item") == 1
    assert ctx.cited == cited_before == ["a21", "b22"]


def test_latex_bullet_without_sources_has_no_cite() -> None:
    out = latex.build_published_claims_section(_NpStore(), _trust([7]), {})
    assert "\\cite" not in out
    assert "\\item ``Nanobuds hold.''; \\url{" in out


def _docx_ctx(store: Any) -> Any:
    return dx._Ctx(
        store=store,
        known_handles=set(),
        doi_links=False,
        library_links=False,
        library_label="x",
        library_search_url="x",
    )


def test_docx_records_sources_and_section_of_claim_cites(monkeypatch) -> None:
    monkeypatch.setattr(dx, "_conjunct_cite_keys", lambda _s, _pk: ["a21", "b22"])
    monkeypatch.setattr(dx, "_render_trust_mark", lambda *_a: None)
    ctx = _docx_ctx(_FindingFake({}))
    ctx.cur_section = "2.1"
    para = docx_lib.Document().add_paragraph()
    dx._render_inline("claim [fi7].", ctx, para)
    assert ctx.claim_cites == {7: ["a21", "b22"]}
    assert ctx.claim_sections == {7: ["2.1"]}


def test_docx_unnumbered_section_is_not_recorded(monkeypatch) -> None:
    monkeypatch.setattr(dx, "_conjunct_cite_keys", lambda _s, _pk: ["a21"])
    monkeypatch.setattr(dx, "_render_trust_mark", lambda *_a: None)
    ctx = _docx_ctx(_FindingFake({}))
    para = docx_lib.Document().add_paragraph()
    dx._render_inline("claim [fi7].", ctx, para)
    assert ctx.claim_sections == {}


def test_docx_published_table_row_per_claim_and_no_new_citations() -> None:
    ctx = _docx_ctx(_NpStore())
    ctx.trust = _trust([7, 8])
    ctx.cited = ["x1", "a21", "b22"]
    ctx.claim_cites = {7: ["b22", "a21"], 8: ["ghost"]}  # ghost never cited
    ctx.claim_sections = {7: ["2", "3.1"]}
    cited_before = list(ctx.cited)
    doc = docx_lib.Document()
    dx._append_published_claims(doc, ctx)

    assert [p.text for p in doc.paragraphs] == [SECTION_TITLE]
    (table,) = doc.tables
    rows = [[c.text for c in r.cells] for r in table.rows]
    assert rows[0] == ["Claim", "Sources", "Cited in", "Artifact"]
    assert len(rows) == 3  # header + one row per claim
    assert rows[1][0] == "“Nanobuds hold.”"
    assert rows[1][1] == "[2, 3]"
    assert rows[1][2] == "2, 3.1"
    assert rows[1][3].startswith("https://w3id.org/np/RAabc — published 2026-01-02")
    assert rows[2][1] == "" and rows[2][2] == ""
    assert ctx.cited == cited_before  # the section registered nothing


# ── doi: prose survives the ref branch (gr493193) ──────────────────────

_DOI_TEXT = "see doi:10.1002/j.2168-0159.2014.tb00055.x now"


def test_docx_doi_prefix_survives() -> None:
    ctx = _docx_ctx(_FindingFake({}))
    para = docx_lib.Document().add_paragraph()
    dx._render_inline(_DOI_TEXT, ctx, para)
    assert para.text == _DOI_TEXT


def test_docx_paper_ref_still_cites() -> None:
    ctx = _docx_ctx(_FindingFake({}))
    para = docx_lib.Document().add_paragraph()
    dx._render_inline("as shown in paper:smith20 here", ctx, para)
    assert ctx.cited == ["smith20"]
    assert "paper:smith20" not in para.text


def test_latex_doi_prefix_survives() -> None:
    ctx = latex._Ctx(keymap={}, known_handles=set())
    out = latex._render_inline(_DOI_TEXT, ctx)
    assert "doi:10.1002/j.2168-0159.2014.tb00055.x" in out


def test_latex_paper_ref_still_cites() -> None:
    ctx = latex._Ctx(keymap={}, known_handles=set())
    out = latex._render_inline("as shown in paper:smith20 here", ctx)
    assert "\\cite{smith20}" in out
    assert ctx.cited == ["smith20"]


def test_docx_known_kinds_stay_dropped_and_doi_survives() -> None:
    ctx = _docx_ctx(_FindingFake({}))
    para = docx_lib.Document().add_paragraph()
    dx._render_inline("see plan:5 and memory:12 here, doi:10.1002/x", ctx, para)
    assert "plan:5" not in para.text and "memory:12" not in para.text
    assert "doi:10.1002/x" in para.text


def test_latex_known_kinds_stay_dropped_and_doi_survives() -> None:
    ctx = latex._Ctx(keymap={}, known_handles=set())
    out = latex._render_inline("see plan:5 and memory:12 here, doi:10.1002/x", ctx)
    assert "plan:5" not in out and "memory:12" not in out
    assert "doi:10.1002/x" in out
